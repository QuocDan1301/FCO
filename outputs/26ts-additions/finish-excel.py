from pathlib import Path
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
import hashlib
import json
import re
import xml.etree.ElementTree as ET
import openpyxl

work = Path(__file__).resolve().parent
target = work.parent.parent / 'gia-cau-thu-bp.xlsx'
original_bytes = target.read_bytes()
original_hash = hashlib.sha256(original_bytes).hexdigest()
before = openpyxl.load_workbook(BytesIO(original_bytes))
sheet = before.worksheets[0]
assert sheet.max_row == 251
assert not any(cell.data_type == 'f' for ws in before for row in ws for cell in row)
cards = json.loads((work / 'cards.json').read_text(encoding='utf-8'))
ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
tag = lambda name: '{' + ns['s'] + '}' + name
with ZipFile(work / 'excel-authored.xlsx') as authored:
    strings = []
    if 'xl/sharedStrings.xml' in authored.namelist():
        for element in ET.fromstring(authored.read('xl/sharedStrings.xml')):
            strings.append(''.join(element.itertext()))
    fresh = ET.fromstring(authored.read('xl/worksheets/sheet1.xml'))
    additions = [deepcopy(row) for row in fresh.find('s:sheetData', ns) if int(row.attrib['r']) >= 252]
assert len(additions) == 17
# Retain the original package and existing XML verbatim. Only graft the authored
# new rows, using the workbook's original cell styles and inline new strings.
missing_price_style = next(cell.style_id for row in sheet.iter_rows(min_row=3) for cell in row if cell.column >= 10 and cell.column % 2 == 0 and cell.value is None and cell.fill.fgColor.rgb == 'FFFFF2B3')
for row in additions:
    for cell in row:
        column = openpyxl.utils.cell.coordinate_from_string(cell.attrib['r'])[0]
        column_number = openpyxl.utils.column_index_from_string(column)
        donor = sheet.cell(251, column_number)
        cell.attrib['s'] = str(missing_price_style if column_number >= 10 and column_number % 2 == 0 else donor.style_id)
        if cell.attrib.get('t') == 's':
            text = strings[int(cell.find('s:v', ns).text)]
            for child in list(cell): cell.remove(child)
            cell.attrib['t'] = 'inlineStr'
            ET.SubElement(ET.SubElement(cell, tag('is')), tag('t')).text = text
        # Blank prices must still have a styled editable cell.
    existing_columns = {openpyxl.utils.cell.coordinate_from_string(c.attrib['r'])[0] for c in row}
    for column_number in range(10, 35, 2):
        column = openpyxl.utils.get_column_letter(column_number)
        if column not in existing_columns:
            ET.SubElement(row, tag('c'), {'r': f"{column}{row.attrib['r']}", 's': str(missing_price_style)})
    row[:] = sorted(row, key=lambda c: openpyxl.utils.column_index_from_string(openpyxl.utils.cell.coordinate_from_string(c.attrib['r'])[0]))
with ZipFile(BytesIO(original_bytes)) as original:
    main = original.read('xl/worksheets/sheet1.xml').decode('utf-8')
    main = main.replace('</sheetData>', ''.join(ET.tostring(row, encoding='unicode') for row in additions) + '</sheetData>')
    main = main.replace('ref="A1:AH251"', 'ref="A1:AH268"').replace('ref="A2:AH251"', 'ref="A2:AH268"')
    missing = sum(sheet.cell(r, c).value is None for r in range(3, 252) for c in range(10, 35, 2)) + 17 * 13
    guide_sheet = before.worksheets[1]
    note = next(cell for row in guide_sheet for cell in row if isinstance(cell.value, str) and re.search(r'\(\d+ ô\)', cell.value))
    guide = original.read('xl/worksheets/sheet2.xml').decode('utf-8')
    updated_note = re.sub(r'\(\d+ ô\)', f'({missing} ô)', note.value)
    replacement = ET.Element(tag('c'), {'r': note.coordinate, 's': str(note.style_id), 't': 'inlineStr'})
    ET.SubElement(ET.SubElement(replacement, tag('is')), tag('t')).text = updated_note
    guide, count = re.subn(r'<c\b[^>]*\br="' + note.coordinate + r'"[^>]*>.*?</c>', ET.tostring(replacement, encoding='unicode'), guide, count=1, flags=re.S)
    assert count == 1
    result = BytesIO()
    with ZipFile(result, 'w') as output:
        for info in original.infolist():
            content = original.read(info.filename)
            if info.filename == 'xl/worksheets/sheet1.xml': content = main.encode('utf-8')
            if info.filename == 'xl/worksheets/sheet2.xml': content = guide.encode('utf-8')
            output.writestr(info, content)
candidate = result.getvalue()
after = openpyxl.load_workbook(BytesIO(candidate))
assert after.sheetnames == before.sheetnames
for old_sheet, new_sheet in zip(before, after):
    assert old_sheet.freeze_panes == new_sheet.freeze_panes
    for row in old_sheet:
        for old in row:
            new = new_sheet[old.coordinate]
            expected = updated_note if old_sheet.title == guide_sheet.title and old.coordinate == note.coordinate else old.value
            assert new.value == expected, (old_sheet.title, old.coordinate)
            assert new._style == old._style, (old_sheet.title, old.coordinate, 'style')
assert after.worksheets[0].auto_filter.ref == 'A2:AH268'
bonuses = [0, 1, 2, 4, 6, 8, 11, 15, 17, 19, 21, 24, 27]
for number, card in enumerate(cards, 252):
    row = after.worksheets[0][number]
    assert row[0].value == card['id'] and row[1].value == card['name']
    for level, bonus in enumerate(bonuses):
        assert row[8 + level * 2].value == card['ovr'] + bonus
        assert row[9 + level * 2].value is None
assert hashlib.sha256(target.read_bytes()).hexdigest() == original_hash, 'Workbook changed during update.'
backup = work / 'excel-before-update.bak'
if not backup.exists(): backup.write_bytes(original_bytes)
temporary = work / 'excel-verified.tmp'
temporary.write_bytes(candidate)
temporary.replace(target)
(work / 'excel-verification.json').write_text(json.dumps({'players': 266, 'added': 17, 'originalCellsAndStylesPreserved': True, 'missingNewPrices': 221, 'filter': 'A2:AH268'}, indent=2))
print('Updated gia-cau-thu-bp.xlsx: 266 players; all original cell values and styles preserved; 17 new rows with +1 to +13 OVR and blank prices.')
